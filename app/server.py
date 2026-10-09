import csv, hashlib, io, json, os, secrets, sqlite3, threading, time, zipfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal, ROUND_CEILING
from functools import wraps
from pathlib import Path
from urllib.parse import urlsplit
from flask import Flask, request, session, jsonify, g, send_file, render_template, abort, current_app
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from itsdangerous import URLSafeTimedSerializer, BadSignature
from PIL import Image
from pypdf import PdfReader
from openpyxl import load_workbook
from .processing import extract, money, date_iso
from .persistence import Database, psycopg
from .storage import FileStore


def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')
def j(value): return json.dumps(value,ensure_ascii=False)
def data():
    d = request.get_json(silent=True)
    if not isinstance(d,dict): raise ValueError('Pedido JSON inválido.')
    return d

def required(d,key,maxlen=1000):
    if not isinstance(d.get(key),str): raise ValueError(f'Campo inválido: {key}.')
    v = d.get(key,'').strip()
    if not v or len(v)>maxlen: raise ValueError(f'Campo obrigatório ou demasiado longo: {key}.')
    return v

def email_address(value):
    import re
    v = str(value).strip().lower()
    if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',v) or len(v)>254: raise ValueError('Email inválido.')
    return v

def db():
    if 'db' not in g:
        g.db = current_app.extensions['database'].connect()
    return g.db

def rows(sql,args=()): return [dict(r) for r in db().execute(sql,args).fetchall()]
def one(sql,args=()):
    r = db().execute(sql,args).fetchone()
    return dict(r) if r else None

def audit(action,resource,ident=None,detail=None):
    db().execute('INSERT INTO audit(user_id,action,resource,resource_id,detail,created) VALUES(?,?,?,?,?,?)',(g.user['id'] if g.user else None,action,resource,ident,j(detail or {}),now()))

def auth(admin=False):
    def decorator(fn):
        @wraps(fn)
        def wrapped(*a,**kw):
            if not g.user: abort(401)
            if admin and g.user['role']!='admin': abort(403)
            return fn(*a,**kw)
        return wrapped
    return decorator

def entity_access(ident,bank=False):
    if not one('SELECT id FROM entities WHERE id=?',(ident,)): abort(404)
    if g.user['role']=='admin': return
    grant = one('SELECT * FROM grants WHERE user_id=? AND entity_id=?',(g.user['id'],ident))
    if not grant or (bank and not grant['bank_access']): abort(403)

def doc_access(ident,include_deleted=False):
    doc = one('SELECT * FROM documents WHERE id=?',(ident,))
    if not doc or (doc['deleted'] and not include_deleted): abort(404)
    if g.user['role']=='admin': return doc
    if doc['visibility']!='financial' or doc['deleted']: abort(403)
    grant = one('SELECT * FROM grants WHERE user_id=? AND entity_id=?',(g.user['id'],doc['entity_id']))
    if not grant: abort(403)
    if not grant['all_documents'] and not one('SELECT 1 FROM grant_documents WHERE grant_id=? AND document_id=?',(grant['id'],ident)): abort(403)
    return doc

def accessible_docs(include_deleted=False):
    if g.user['role']=='admin':
        return rows('SELECT * FROM documents'+('' if include_deleted else ' WHERE deleted IS NULL')+' ORDER BY id DESC')
    return rows('''SELECT DISTINCT d.* FROM documents d JOIN grants gr ON gr.entity_id=d.entity_id
        WHERE gr.user_id=? AND d.visibility='financial' AND d.deleted IS NULL
        AND (gr.all_documents=1 OR EXISTS(SELECT 1 FROM grant_documents gd WHERE gd.grant_id=gr.id AND gd.document_id=d.id)) ORDER BY d.id DESC''',(g.user['id'],))

def paid_cents(doc_id): return one('SELECT coalesce(sum(amount),0) AS amount FROM allocations WHERE document_id=?',(doc_id,))['amount']

def doc_out(doc):
    d = {k:v for k,v in doc.items() if k not in ('storage','sha256')}
    d['metadata'] = json.loads(d['metadata']); d['extraction'] = json.loads(d['extraction'])
    d['paid_cents'] = paid_cents(d['id'])
    total = money(d['metadata'].get('total'))
    d['payment'] = 'pago' if total is not None and total > 0 and d['paid_cents']>=total else 'parcialmente_pago' if d['paid_cents']>0 else 'por_pagar'
    if g.user['role']!='admin': d.pop('work_id',None)
    return d

def private_record(kind,ident):
    r = one('SELECT * FROM private_records WHERE kind=? AND id=?',(kind,ident))
    if not r: abort(404)
    r['data'] = json.loads(r['data'])
    return r


def create_app(config=None):
    app = Flask(__name__)
    root = Path(os.environ.get('APP_DATA_DIR',str(Path(__file__).resolve().parent.parent/'instance')))
    if config and config.get('DATA_DIR'): root = Path(config['DATA_DIR'])
    root.mkdir(parents=True,exist_ok=True,mode=0o700)
    os.chmod(root,0o700)
    (root/'files').mkdir(exist_ok=True,mode=0o700)
    keyfile = root/'session.key'
    if not keyfile.exists():
        try:
            with keyfile.open('x') as f: f.write(secrets.token_hex(64))
            os.chmod(keyfile,0o600)
        except FileExistsError: pass
    app.config.update(SECRET_KEY=os.environ.get('APP_SECRET_KEY') or keyfile.read_text().strip(),MAX_CONTENT_LENGTH=25*1024*1024,
        SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Lax',SESSION_COOKIE_SECURE=os.environ.get('APP_HTTPS')=='1',
        PERMANENT_SESSION_LIFETIME=timedelta(hours=12),DATA_DIR=root,DATABASE=root/'app.sqlite3',TESTING=False,START_WORKER=os.environ.get('APP_WORKER','1')=='1',
        DATABASE_URL=os.environ.get('DATABASE_URL'),DB_SCHEMA=os.environ.get('APP_DB_SCHEMA','conta_clara'),
        STORAGE_BACKEND=os.environ.get('APP_STORAGE_BACKEND','local'),SUPABASE_URL=os.environ.get('SUPABASE_URL'),
        SUPABASE_STORAGE_KEY=os.environ.get('SUPABASE_STORAGE_KEY'),SUPABASE_BUCKET=os.environ.get('SUPABASE_BUCKET','conta-clara'),
        SETUP_TOKEN=os.environ.get('APP_SETUP_TOKEN'),REQUIRE_EXTERNAL_STORAGE=os.environ.get('APP_REQUIRE_EXTERNAL_STORAGE')=='1')
    if config: app.config.update(config)
    if app.config['REQUIRE_EXTERNAL_STORAGE'] and not (app.config['DATABASE_URL'] and app.config['STORAGE_BACKEND']=='supabase' and os.environ.get('APP_SECRET_KEY') and app.config['SETUP_TOKEN']):
        raise RuntimeError('O alojamento exige PostgreSQL, armazenamento Supabase privado, APP_SECRET_KEY e APP_SETUP_TOKEN. Não será usada persistência local efémera.')
    database=Database(app.config); app.extensions['database']=database
    database.initialize(app.config['START_WORKER'])
    store=FileStore(app.config); app.extensions['file_store']=store
    store.initialize()
    if store.remote: database.secure_storage_policy(store.bucket)
    signer = URLSafeTimedSerializer(app.secret_key,salt='private-download')

    @app.before_request
    def before():
        g.app_db = app.config['DATABASE']; g.user = None
        if session.get('user_id'):
            g.user = one('SELECT id,email,name,role FROM users WHERE id=? AND active=1',(session['user_id'],))
        if request.path.startswith('/api/') and request.method not in ('GET','HEAD','OPTIONS'):
            token = request.headers.get('X-CSRF-Token','')
            if not session.get('csrf') or not secrets.compare_digest(session['csrf'],token): abort(403,description='Sessão expirada. Atualize a página.')

    @app.teardown_appcontext
    def teardown(error):
        con = g.pop('db',None)
        if con is not None: con.close()

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['Referrer-Policy']='no-referrer'
        response.headers['Content-Security-Policy']="default-src 'self'; img-src 'self' blob: data:; style-src 'self'; script-src 'self'; frame-src 'self' blob:; object-src 'none'; base-uri 'self'; frame-ancestors 'self'; form-action 'self'"
        response.headers['Cache-Control']='no-store' if request.path.startswith(('/api/','/download/')) else 'no-cache'
        if app.config['SESSION_COOKIE_SECURE']: response.headers['Strict-Transport-Security']='max-age=31536000'
        return response

    @app.errorhandler(ValueError)
    def invalid(error): return jsonify(error=str(error)),400
    @app.errorhandler(sqlite3.IntegrityError)
    def conflict(error): return jsonify(error='Este registo já existe ou tem uma associação inválida.'),409
    app.register_error_handler(psycopg.IntegrityError,conflict)
    @app.errorhandler(__import__('decimal').InvalidOperation)
    def invalid_decimal(error): return jsonify(error='Número inválido.'),400
    @app.errorhandler(413)
    def large(error): return jsonify(error='Limite de 25 MB por pedido. Carregue os ficheiros separadamente.'),413
    for code in (400,401,403,404,409,429):
        app.register_error_handler(code,lambda error: (jsonify(error=error.description if error.code in (400,409,429) else {401:'Autenticação necessária.',403:'Acesso não autorizado.',404:'Registo não encontrado.'}.get(error.code,'Pedido inválido.')),error.code))

    @app.get('/')
    def index(): return render_template('index.html')
    @app.get('/health')
    def health():
        db().execute('SELECT 1')
        return jsonify(status='ok')
    @app.get('/api/session')
    def session_info():
        if not session.get('csrf'): session['csrf']=secrets.token_urlsafe(32)
        return jsonify(user=g.user,csrf=session['csrf'],needs_setup=not bool(one('SELECT id FROM users LIMIT 1')),setup_requires_token=bool(app.config['SETUP_TOKEN']))

    def password(d):
        v = required(d,'password',256)
        if len(v)<12: raise ValueError('A palavra-passe deve ter pelo menos 12 caracteres.')
        return v
    def create_user(d,role):
        return db().execute('INSERT INTO users(email,name,password,role,created) VALUES(?,?,?,?,?)',(email_address(required(d,'email')),required(d,'name',120),generate_password_hash(password(d)),role,now())).lastrowid
    def logged(ident):
        session.clear(); session['user_id']=ident; session['csrf']=secrets.token_urlsafe(32); session.permanent=True
        return jsonify(user=one('SELECT id,email,name,role FROM users WHERE id=?',(ident,)),csrf=session['csrf'])

    @app.post('/api/setup')
    def setup():
        d=data()
        if app.config['SETUP_TOKEN'] and not secrets.compare_digest(str(d.get('setup_token','')),app.config['SETUP_TOKEN']): abort(403)
        db().execute('BEGIN IMMEDIATE')
        if one('SELECT id FROM users LIMIT 1'): abort(409,description='O administrador já foi criado.')
        ident=create_user(d,'admin'); audit('admin_created','users',ident); db().commit()
        return logged(ident)

    def invite(token):
        i=one('SELECT * FROM invitations WHERE token_hash=?',(hashlib.sha256(str(token).encode()).hexdigest(),))
        if not i or i['used'] or i['expires']<now(): raise ValueError('Convite inválido, expirado ou já utilizado.')
        return i
    @app.post('/api/register-invite')
    def register_invite():
        d=data(); i=invite(required(d,'token'))
        if email_address(required(d,'email'))!=i['email']: abort(403)
        if one('SELECT id FROM users WHERE email=?',(i['email'],)): abort(409,description='Este email já tem conta. Inicie sessão para aceitar o convite.')
        ident=create_user(d,'accountant'); audit('account_created','users',ident); db().commit()
        return logged(ident)
    @app.post('/api/login')
    def login():
        d=data(); email=email_address(required(d,'email')); key=hashlib.sha256((email+'|'+request.remote_addr).encode()).hexdigest()
        attempt=one('SELECT * FROM login_attempts WHERE key=?',(key,))
        if attempt and attempt['count']>=8 and time.time()-attempt['last']<900: abort(429,description='Demasiadas tentativas. Aguarde 15 minutos.')
        user=one('SELECT * FROM users WHERE email=? AND active=1',(email,))
        if not user or not check_password_hash(user['password'],str(d.get('password',''))):
            count=attempt['count']+1 if attempt and time.time()-attempt['last']<900 else 1
            db().execute('INSERT OR REPLACE INTO login_attempts VALUES(?,?,?)',(key,count,time.time())); db().commit(); abort(401)
        db().execute('DELETE FROM login_attempts WHERE key=?',(key,)); audit('login','users',user['id']); db().commit()
        return logged(user['id'])
    @app.post('/api/logout')
    @auth()
    def logout(): session.clear(); return jsonify(ok=True)
    @app.post('/api/password')
    @auth()
    def change_password():
        d=data(); user=one('SELECT password FROM users WHERE id=?',(g.user['id'],))
        if not check_password_hash(user['password'],str(d.get('current_password',''))): raise ValueError('Palavra-passe atual incorreta.')
        db().execute('UPDATE users SET password=? WHERE id=?',(generate_password_hash(password(d)),g.user['id'])); audit('password_changed','users',g.user['id']); db().commit(); session.clear()
        return jsonify(ok=True)

    @app.get('/api/entities')
    @auth()
    def entities():
        return jsonify(rows('SELECT * FROM entities ORDER BY name') if g.user['role']=='admin' else rows('SELECT e.* FROM entities e JOIN grants gr ON gr.entity_id=e.id WHERE gr.user_id=? ORDER BY e.name',(g.user['id'],)))
    @app.post('/api/entities')
    @auth(True)
    def add_entity():
        d=data(); name=required(d,'name',150); nif=required(d,'nif',9)
        if not nif.isdigit() or len(nif)!=9: raise ValueError('O NIF deve ter 9 algarismos.')
        if d.get('kind') not in ('pessoal','empresa'): raise ValueError('Tipo de entidade inválido.')
        ident=db().execute('INSERT INTO entities(name,nif,kind) VALUES(?,?,?)',(name,nif,d['kind'])).lastrowid
        audit('created','entities',ident); db().commit(); return jsonify(id=ident),201

    def scope_validate(scopes):
        if not isinstance(scopes,list) or not scopes: raise ValueError('Selecione pelo menos uma entidade.')
        result=[]
        for s in scopes:
            eid=int(s['entity_id']); entity_access(eid)
            ids=list(set(int(x) for x in s.get('document_ids',[])))
            for ident in ids:
                doc=doc_access(ident)
                if doc['entity_id']!=eid or doc['visibility']!='financial': raise ValueError('Documento privado ou de outra entidade.')
            result.append({'entity_id':eid,'all_documents':bool(s.get('all_documents')),'bank_access':bool(s.get('bank_access')),'document_ids':ids})
        if len({s['entity_id'] for s in result})!=len(result): raise ValueError('Entidades repetidas.')
        return result
    def apply_scope(user_id,scopes):
        db().execute('DELETE FROM grants WHERE user_id=?',(user_id,))
        for s in scopes:
            gid=db().execute('INSERT INTO grants(user_id,entity_id,all_documents,bank_access) VALUES(?,?,?,?)',(user_id,s['entity_id'],int(s['all_documents']),int(s['bank_access']))).lastrowid
            for ident in s['document_ids']: db().execute('INSERT INTO grant_documents VALUES(?,?)',(gid,ident))
    @app.post('/api/invitations')
    @auth(True)
    def add_invite():
        d=data(); scope=scope_validate(d.get('scope')); email=email_address(required(d,'email')); token=secrets.token_urlsafe(40)
        ident=db().execute('INSERT INTO invitations(email,token_hash,scope,expires,created) VALUES(?,?,?,?,?)',(email,hashlib.sha256(token.encode()).hexdigest(),j(scope),(datetime.now(timezone.utc)+timedelta(days=7)).isoformat(timespec='seconds'),now())).lastrowid
        audit('invited','invitations',ident,{'email':email}); db().commit()
        return jsonify(id=ident,path='/?invite='+token,expires_days=7),201
    @app.post('/api/invitations/accept')
    @auth()
    def accept_invite():
        d=data(); db().execute('BEGIN IMMEDIATE'); i=invite(required(d,'token'))
        if g.user['email']!=i['email'] or g.user['role']!='accountant': abort(403)
        scopes=json.loads(i['scope'])
        # Recheck stored selections: a document may have become private since invitation.
        for s in scopes:
            for ident in s['document_ids']:
                doc=one('SELECT entity_id,visibility,deleted FROM documents WHERE id=?',(ident,))
                if not doc or doc['deleted'] or doc['visibility']!='financial' or doc['entity_id']!=s['entity_id']: raise ValueError('O convite precisa de ser atualizado pelo administrador.')
        apply_scope(g.user['id'],scopes); db().execute('UPDATE invitations SET used=? WHERE id=?',(now(),i['id'])); audit('accepted','invitations',i['id']); db().commit()
        return jsonify(ok=True)
    @app.get('/api/access')
    @auth(True)
    def access_list():
        accounts=rows("SELECT id,email,name,active FROM users WHERE role='accountant'")
        for account in accounts:
            account['grants']=rows('SELECT * FROM grants WHERE user_id=?',(account['id'],))
            for grant in account['grants']: grant['document_ids']=[r['document_id'] for r in rows('SELECT document_id FROM grant_documents WHERE grant_id=?',(grant['id'],))]
        return jsonify(accounts=accounts,invitations=rows('SELECT id,email,expires,used FROM invitations ORDER BY id DESC'))
    @app.put('/api/access/<int:ident>')
    @auth(True)
    def change_access(ident):
        d=data(); db().execute('BEGIN IMMEDIATE')
        if not one("SELECT id FROM users WHERE id=? AND role='accountant'",(ident,)): abort(404)
        scopes=scope_validate(d['scope']) if d.get('scope') else []
        apply_scope(ident,scopes)
        # Old pending invitations must not restore permissions after revocation.
        email=one('SELECT email FROM users WHERE id=?',(ident,))['email']
        db().execute('UPDATE invitations SET used=? WHERE email=? AND used IS NULL',(now(),email))
        audit('scope_changed','users',ident); db().commit(); return jsonify(ok=True)
    @app.delete('/api/invitations/<int:ident>')
    @auth(True)
    def revoke_invite(ident):
        db().execute('UPDATE invitations SET used=? WHERE id=?',(now(),ident)); audit('revoked','invitations',ident); db().commit(); return jsonify(ok=True)

    def save_upload(file):
        raw=file.read(20*1024*1024+1)
        if not raw or len(raw)>20*1024*1024: raise ValueError('Cada ficheiro deve ter entre 1 byte e 20 MB.')
        suffix=Path(file.filename or '').suffix.lower()
        if suffix=='.pdf' and raw.startswith(b'%PDF-'):
            try:
                reader=PdfReader(io.BytesIO(raw))
                if reader.is_encrypted or not reader.pages: raise ValueError('PDF encriptado ou vazio.')
                if len(reader.pages)>100: raise ValueError('Limite de 100 páginas por documento.')
            except ValueError: raise
            except Exception as e: raise ValueError('PDF inválido.') from e
            mime='application/pdf'
        elif suffix in ('.jpg','.jpeg','.png'):
            try:
                im=Image.open(io.BytesIO(raw)); im.verify()
                if im.format not in ('PNG','JPEG'): raise ValueError('Imagem inválida.')
                if im.width*im.height>40000000: raise ValueError('Limite de 40 megapíxeis por imagem.')
                mime='image/png' if im.format=='PNG' else 'image/jpeg'
            except Exception as e: raise ValueError('Imagem inválida.') from e
        else: raise ValueError('Aceites apenas PDF, JPG e PNG.')
        return raw,mime,secure_filename(file.filename) or 'documento'+suffix

    @app.post('/api/documents')
    @auth(True)
    def upload():
        eid=int(request.form.get('entity_id',0)); entity_access(eid)
        visibility=request.form.get('visibility','financial')
        if visibility not in ('financial','private'): raise ValueError('Visibilidade inválida.')
        file=request.files.get('file')
        if not file: raise ValueError('Escolha um ficheiro.')
        raw,mime,filename=save_upload(file); sha=hashlib.sha256(raw).hexdigest()
        db().execute('BEGIN IMMEDIATE')
        duplicates=rows('SELECT id,filename FROM documents WHERE sha256=? AND deleted IS NULL',(sha,))
        if duplicates and request.form.get('confirm_duplicate')!='true': return jsonify(error='Possível ficheiro duplicado. Confirme antes de continuar.',duplicates=duplicates),409
        stored=secrets.token_hex(24)+Path(filename).suffix.lower(); path=root/'files'/stored
        try:
            store.put(stored,raw,mime)
            ident=db().execute('INSERT INTO documents(entity_id,filename,storage,mime,sha256,created,visibility) VALUES(?,?,?,?,?,?,?)',(eid,filename,stored,mime,sha,now(),visibility)).lastrowid
            audit('uploaded','documents',ident); db().commit()
        except Exception:
            if path.exists(): store.discard_failed_upload(stored)
            raise
        return jsonify(document=doc_out(doc_access(ident))),201
    @app.get('/api/documents')
    @auth()
    def documents():
        docs=accessible_docs(request.args.get('deleted')=='1' and g.user['role']=='admin')
        return jsonify([doc_out(d) for d in docs])
    @app.get('/api/documents/<int:ident>')
    @auth()
    def document(ident):
        out=doc_out(doc_access(ident)); out['comments']=rows('SELECT c.*,u.name AS author FROM comments c JOIN users u ON u.id=c.user_id WHERE document_id=? ORDER BY c.id',(ident,))
        return jsonify(out)
    @app.put('/api/documents/<int:ident>')
    @auth(True)
    def update_document(ident):
        db().execute('BEGIN IMMEDIATE'); old=doc_access(ident); d=data(); meta=d.get('metadata',{})
        allowed={'type','supplier','supplier_nif','recipient','recipient_nif','number','date','due_date','net','vat','vat_rates','total','currency','description','category'}
        if not isinstance(meta,dict) or set(meta)-allowed: raise ValueError('Campos de documento inválidos.')
        for key in ('date','due_date'):
            if meta.get(key): meta[key]=date_iso(meta[key])
        for key in ('net','vat','total'):
            if meta.get(key) not in (None,''):
                value=money(meta[key])
                if value<0: raise ValueError('Indique valores positivos; use o tipo Nota de crédito para créditos.')
                meta[key]=value/100
        for key,value in meta.items():
            if isinstance(value,str) and len(value)>3000: raise ValueError('Campo demasiado longo.')
        if meta.get('currency') and meta['currency']!='EUR': raise ValueError('Esta versão suporta apenas EUR; confirme a conversão manualmente.')
        if money(meta.get('total')) is not None and paid_cents(ident)>money(meta['total']): raise ValueError('O total não pode ser inferior aos pagamentos associados.')
        if paid_cents(ident) and money(meta.get('total')) is None: raise ValueError('Não pode remover o total de um documento com pagamentos associados.')
        if paid_cents(ident) and meta.get('type')!=json.loads(old['metadata']).get('type'): raise ValueError('Remova os pagamentos associados antes de mudar o tipo de documento.')
        eid=int(d.get('entity_id',old['entity_id'])); entity_access(eid)
        if eid!=old['entity_id'] and one('SELECT id FROM allocations WHERE document_id=?',(ident,)): raise ValueError('Remova as associações antes de mudar a entidade.')
        review=d.get('review',old['review'])
        if review not in ('por_rever','validado'): raise ValueError('Estado de revisão inválido.')
        if review=='validado' and (not meta.get('date') or not meta.get('number') or money(meta.get('total')) is None): raise ValueError('Para validar, confirme a data, o número e o total.')
        work=d.get('work_id') or None
        if work: private_record('works',int(work))
        duplicate=[]
        if all(meta.get(k) not in ('',None) for k in ('supplier_nif','number','date','total')):
            for other in rows('SELECT id,filename,metadata FROM documents WHERE id!=? AND entity_id=? AND deleted IS NULL',(ident,eid)):
                om=json.loads(other['metadata'])
                if all(str(om.get(k))==str(meta.get(k)) for k in ('supplier_nif','number','date','total')): duplicate.append({'id':other['id'],'filename':other['filename']})
        if duplicate and not d.get('confirm_duplicate'): return jsonify(error='Possível documento duplicado por NIF, número, data e total.',duplicates=duplicate),409
        visibility=d.get('visibility',old['visibility'])
        if visibility not in ('financial','private'): raise ValueError('Visibilidade inválida.')
        db().execute('UPDATE documents SET metadata=?,review=?,entity_id=?,work_id=?,visibility=? WHERE id=?',(j(meta),review,eid,work,visibility,ident))
        if eid!=old['entity_id'] or visibility=='private': db().execute('DELETE FROM grant_documents WHERE document_id=?',(ident,))
        audit('updated','documents',ident,{'review_before':old['review'],'review_after':review}); db().commit(); return jsonify(document=doc_out(doc_access(ident)))
    @app.delete('/api/documents/<int:ident>')
    @auth(True)
    def delete_document(ident):
        doc_access(ident); db().execute('UPDATE documents SET deleted=? WHERE id=?',(now(),ident)); audit('soft_deleted','documents',ident); db().commit(); return jsonify(ok=True)
    @app.post('/api/documents/<int:ident>/restore')
    @auth(True)
    def restore_document(ident):
        doc_access(ident,True); db().execute('UPDATE documents SET deleted=NULL WHERE id=?',(ident,)); audit('restored','documents',ident); db().commit(); return jsonify(ok=True)
    @app.post('/api/documents/<int:ident>/retry')
    @auth(True)
    def retry_document(ident):
        doc=doc_access(ident)
        if doc['ocr_status']=='processing': abort(409,description='O processamento ainda está em curso.')
        db().execute("UPDATE documents SET ocr_status='queued',ocr_error=NULL WHERE id=?",(ident,)); audit('ocr_retried','documents',ident); db().commit(); return jsonify(ok=True)
    @app.post('/api/documents/<int:ident>/accounting')
    @auth()
    def accounting_state(ident):
        doc=doc_access(ident); d=data(); state=d.get('state')
        if state not in ('por_tratar','em_esclarecimento','tratado'): raise ValueError('Estado inválido.')
        db().execute('UPDATE documents SET accounting=? WHERE id=?',(state,ident)); audit('accounting_state','documents',ident,{'before':doc['accounting'],'after':state}); db().commit(); return jsonify(ok=True)
    @app.post('/api/documents/<int:ident>/comments')
    @auth()
    def comment(ident):
        doc_access(ident); d=data(); db().execute('INSERT INTO comments(document_id,user_id,body,created) VALUES(?,?,?,?)',(ident,g.user['id'],required(d,'body',5000),now())); audit('commented','documents',ident); db().commit(); return jsonify(ok=True),201
    @app.get('/api/documents/<int:ident>/link')
    @auth()
    def download_link(ident):
        doc_access(ident); token=signer.dumps({'document':ident,'user':g.user['id']})
        return jsonify(url='/download/'+token,expires_seconds=300)
    @app.get('/download/<token>')
    @auth()
    def download(token):
        try: payload=signer.loads(token,max_age=300)
        except BadSignature: abort(403)
        if payload.get('user')!=g.user['id']: abort(403)
        doc=doc_access(payload['document'])
        audit('downloaded','documents',doc['id']); db().commit()
        return send_file(store.get_path(doc['storage'],doc['sha256']),mimetype=doc['mime'],download_name=doc['filename'],as_attachment=request.args.get('inline')!='1',conditional=True)

    @app.get('/api/movements')
    @auth()
    def movements():
        ms=rows('SELECT * FROM movements ORDER BY date DESC,id DESC') if g.user['role']=='admin' else rows('SELECT m.* FROM movements m JOIN grants gr ON gr.entity_id=m.entity_id WHERE gr.user_id=? AND gr.bank_access=1 ORDER BY m.date DESC,m.id DESC',(g.user['id'],))
        for m in ms:
            m['allocated_cents']=one('SELECT coalesce(sum(amount),0) amount FROM allocations WHERE movement_id=?',(m['id'],))['amount']
            if g.user['role']!='admin' and m['source_id']:
                permitted={d['id'] for d in accessible_docs()}
                if m['source_id'] not in permitted: m['source_id']=None
        return jsonify(ms)
    @app.post('/api/bank/preview')
    @auth(True)
    def bank_preview():
        f=request.files.get('file')
        if not f: raise ValueError('Escolha o extrato.')
        raw=f.read(20*1024*1024+1)
        if not raw or len(raw)>20*1024*1024: raise ValueError('Limite de 20 MB.')
        suffix=Path(f.filename).suffix.lower(); out=[]; source_id=None
        if suffix=='.csv':
            try: text=raw.decode('utf-8-sig')
            except UnicodeDecodeError: text=raw.decode('cp1252')
            try: dialect=csv.Sniffer().sniff(text[:8000],delimiters=';,\t|')
            except csv.Error: dialect=csv.excel; dialect.delimiter=';'
            table=list(csv.reader(io.StringIO(text),dialect))
            if not table: raise ValueError('Extrato vazio.')
            if len(set(table[0]))!=len(table[0]): raise ValueError('Os cabeçalhos devem ser únicos.')
            out=[dict(zip(table[0],r)) for r in table[1:] if any(r)]
        elif suffix=='.xls':
            import xlrd
            try:
                workbook=xlrd.open_workbook(file_contents=raw)
                sheet=workbook.sheet_by_index(0)
                heads=[str(sheet.cell_value(0,i) or f'Coluna {i+1}') for i in range(sheet.ncols)]
                if len(set(heads))!=len(heads): raise ValueError('Os cabeçalhos devem ser únicos.')
                for ri in range(1,sheet.nrows):
                    row={}
                    for ci,key in enumerate(heads):
                        cell=sheet.cell(ri,ci)
                        row[key]=xlrd.xldate_as_datetime(cell.value,workbook.datemode).date().isoformat() if cell.ctype==xlrd.XL_CELL_DATE else cell.value
                    if any(v not in ('',None) for v in row.values()): out.append(row)
            except xlrd.XLRDError as e: raise ValueError('Ficheiro Excel inválido.') from e
        elif suffix=='.xlsx':
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as zipped:
                    if sum(info.file_size for info in zipped.infolist())>100*1024*1024: raise ValueError('Excel demasiado grande após descompressão.')
            except zipfile.BadZipFile as e: raise ValueError('Ficheiro Excel inválido.') from e
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            if workbook.active.max_row>5001 or workbook.active.max_column>100: raise ValueError('Limite de 5000 movimentos e 100 colunas.')
            table=list(workbook.active.iter_rows(values_only=True)); workbook.close()
            if not table: raise ValueError('Extrato vazio.')
            heads=[str(v or f'Coluna {i+1}') for i,v in enumerate(table[0])]
            if len(set(heads))!=len(heads): raise ValueError('Os cabeçalhos devem ser únicos.')
            out=[dict(zip(heads,[v.isoformat()[:10] if isinstance(v,datetime) else v for v in r])) for r in table[1:] if any(v is not None for v in r)]
        elif suffix=='.pdf':
            eid=int(request.form.get('entity_id',0)); entity_access(eid)
            f.stream=io.BytesIO(raw); verified,mime,filename=save_upload(f)
            stored=secrets.token_hex(24)+'.pdf'; path=root/'files'/stored
            store.put(stored,verified,mime)
            try:
                source_id=db().execute("INSERT INTO documents(entity_id,filename,storage,mime,sha256,created,metadata) VALUES(?,?,?,?,?,?,?)",(eid,filename,stored,mime,hashlib.sha256(raw).hexdigest(),now(),j({'type':'Extrato bancário'}))).lastrowid
                audit('bank_source_uploaded','documents',source_id); db().commit()
            except Exception:
                store.discard_failed_upload(stored); raise
            return jsonify(rows=[],columns=[],source_id=source_id,pending=True,warning='O original foi guardado. O PDF está em fila para extração em segundo plano; os movimentos ainda não foram importados.')
        else: raise ValueError('Aceites CSV, Excel .xlsx/.xls e PDF.')
        if len(out)>5000: raise ValueError('Importe no máximo 5000 movimentos de cada vez.')
        if not source_id:
            eid=int(request.form.get('entity_id',0)); entity_access(eid)
            sha=hashlib.sha256(raw).hexdigest()
            existing=one("SELECT id FROM documents WHERE entity_id=? AND sha256=? AND deleted IS NULL AND visibility='financial'",(eid,sha))
            if existing: source_id=existing['id']
            else:
                filename=secure_filename(f.filename) or 'extrato'+suffix
                stored=secrets.token_hex(24)+suffix; path=root/'files'/stored
                mime={'.csv':'text/csv','.xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','.xls':'application/vnd.ms-excel'}[suffix]
                try:
                    store.put(stored,raw,mime)
                    source_id=db().execute("INSERT INTO documents(entity_id,filename,storage,mime,sha256,created,metadata,ocr_status) VALUES(?,?,?,?,?,?,?,'completed')",(eid,filename,stored,mime,sha,now(),j({'type':'Extrato bancário'}))).lastrowid
                    audit('bank_source_uploaded','documents',source_id); db().commit()
                except Exception:
                    if path.exists(): store.discard_failed_upload(stored)
                    raise
        return jsonify(rows=out,columns=list(out[0]) if out else [],source_id=source_id,warning='Confirme todas as linhas e o sinal dos montantes antes de importar.')

    @app.get('/api/bank/preview/<int:ident>')
    @auth(True)
    def bank_pdf_preview(ident):
        import re
        doc=doc_access(ident)
        if doc['mime']!='application/pdf': raise ValueError('Selecione um extrato PDF.')
        if doc['ocr_status'] in ('queued','processing'): return jsonify(pending=True,source_id=ident,rows=[],columns=[])
        if doc['ocr_status']=='error': return jsonify(pending=False,source_id=ident,rows=[],columns=[],warning=doc['ocr_error'])
        out=[]; date_pattern=r'(\d{2}[/-]\d{2}[/-]\d{4}|\d{4}-\d{2}-\d{2})'; amount_pattern=r'(-?\d[\d.]*[,.]\d{2})'
        for line in json.loads(doc['extraction']).get('text','').splitlines():
            cells=[c.strip() for c in re.split(r';|\t|\s{2,}',line) if c.strip()]
            if cells and re.fullmatch(date_pattern,cells[0]) and len(cells)>=3:
                out.append({f'Coluna {i+1}':v for i,v in enumerate(cells)})
                continue
            match=re.fullmatch(r'\s*'+date_pattern+r'\s+(?:'+date_pattern+r'\s+)?(.+?)\s+'+amount_pattern+r'(?:\s+'+amount_pattern+r')?\s*',line)
            if match:
                dt,vd,desc,amount,balance=match.groups()
                out.append({'Data':dt,'Data-valor':vd or '', 'Descrição':desc,'Montante':amount,'Saldo':balance or ''})
        if len(out)>5000: raise ValueError('Limite de 5000 linhas por importação.')
        columns=list(dict.fromkeys(key for row in out for key in row))
        return jsonify(pending=False,rows=out,columns=columns,source_id=ident,warning='Confirme o mapeamento e todas as linhas extraídas do PDF.' if out else 'Não foi possível separar os movimentos. Consulte o original e introduza-os manualmente.')

    def movement_values(d,eid,bank,account,source):
        dt=date_iso(required(d,'date')); vd=date_iso(d.get('value_date')); description=required(d,'description',2000)
        amount=money(d.get('amount')); balance=money(d.get('balance'))
        if amount is None or amount==0: raise ValueError('O montante deve ser diferente de zero. Saídas negativas, entradas positivas.')
        fingerprint=hashlib.sha256(j([eid,bank.casefold(),account.replace(' ','').casefold(),dt,description.casefold(),amount]).encode()).hexdigest()
        return (eid,bank,account,dt,vd,description,amount,balance,source,fingerprint,now())
    @app.post('/api/bank/import')
    @auth(True)
    def bank_import():
        d=data(); db().execute('BEGIN IMMEDIATE'); eid=int(d.get('entity_id',0)); entity_access(eid); bank=required(d,'bank',150); account=required(d,'account',100); source=d.get('source_id') or None
        if source and doc_access(int(source))['entity_id']!=eid: raise ValueError('O extrato deve pertencer à mesma entidade.')
        items=d.get('rows')
        if not isinstance(items,list) or not items or len(items)>5000: raise ValueError('Selecione entre 1 e 5000 movimentos.')
        parsed=[movement_values(r,eid,bank,account,source) for r in items]; duplicates=[]; seen=set()
        for idx,v in enumerate(parsed):
            if v[-2] in seen or one('SELECT id FROM movements WHERE fingerprint=?',(v[-2],)): duplicates.append(idx+1)
            seen.add(v[-2])
        if duplicates and not d.get('confirm_duplicates'): return jsonify(error='Movimentos possivelmente duplicados. Confirme antes de importar.',duplicate_rows=duplicates),409
        for v in parsed:
            ident=db().execute('INSERT INTO movements(entity_id,bank,account,date,value_date,description,amount,balance,source_id,fingerprint,created) VALUES(?,?,?,?,?,?,?,?,?,?,?)',v).lastrowid
            audit('imported','movements',ident)
        if source:
            original=doc_access(int(source)); meta=json.loads(original['metadata'])
            if not meta.get('date'):
                meta['date']=min(v[3] for v in parsed)
                db().execute('UPDATE documents SET metadata=? WHERE id=?',(j(meta),source))
        db().commit(); return jsonify(imported=len(parsed)),201

    def movement_access(ident):
        m=one('SELECT * FROM movements WHERE id=?',(ident,))
        if not m: abort(404)
        entity_access(m['entity_id'],True); return m
    @app.put('/api/movements/<int:ident>')
    @auth(True)
    def classify_movement(ident):
        movement_access(ident); d=data(); state=d.get('classification')
        if state not in ('por_classificar','comissao','transferencia_interna','sem_fatura','pagamento'): raise ValueError('Classificação inválida.')
        db().execute('UPDATE movements SET classification=? WHERE id=?',(state,ident)); audit('classified','movements',ident,{'classification':state}); db().commit(); return jsonify(ok=True)
    @app.get('/api/allocations')
    @auth()
    def allocations():
        allowed={d['id'] for d in accessible_docs()}
        result=[]
        for r in rows('SELECT a.*,m.entity_id FROM allocations a JOIN movements m ON m.id=a.movement_id ORDER BY a.id DESC'):
            if r['document_id'] in allowed:
                if g.user['role']=='admin' or one('SELECT id FROM grants WHERE user_id=? AND entity_id=? AND bank_access=1',(g.user['id'],r['entity_id'])): result.append(r)
        return jsonify(result)
    @app.post('/api/allocations')
    @auth(True)
    def add_allocation():
        d=data(); db().execute('BEGIN IMMEDIATE'); mid=int(d.get('movement_id',0)); did=int(d.get('document_id',0)); m=movement_access(mid); doc=doc_access(did); amount=money(d.get('amount'))
        if m['entity_id']!=doc['entity_id']: raise ValueError('Não pode associar entidades diferentes.')
        meta=json.loads(doc['metadata']); total=money(meta.get('total'))
        if doc['visibility']!='financial' or meta.get('type') not in ('Fatura','Fatura-recibo','Recibo','Nota de crédito'): raise ValueError('Associe apenas faturas, recibos ou notas de crédito financeiros.')
        if (meta.get('type')=='Nota de crédito')!=(m['amount']>0): raise ValueError('Use saídas para pagar faturas e entradas para reembolsos de notas de crédito.')
        used=one('SELECT coalesce(sum(amount),0) amount FROM allocations WHERE movement_id=?',(mid,))['amount']
        if amount is None or amount<=0 or amount>abs(m['amount'])-used or total is None or amount>total-paid_cents(did): raise ValueError('A associação excede o saldo do movimento ou da fatura, ou falta confirmar o total.')
        ident=db().execute('INSERT INTO allocations(movement_id,document_id,amount,created,user_id) VALUES(?,?,?,?,?)',(mid,did,amount,now(),g.user['id'])).lastrowid
        audit('payment_linked','allocations',ident,{'movement_id':mid,'document_id':did,'amount':amount}); db().commit(); return jsonify(id=ident),201
    @app.delete('/api/allocations/<int:ident>')
    @auth(True)
    def remove_allocation(ident):
        if not one('SELECT id FROM allocations WHERE id=?',(ident,)): abort(404)
        db().execute('DELETE FROM allocations WHERE id=?',(ident,)); audit('payment_unlinked','allocations',ident); db().commit(); return jsonify(ok=True)
    @app.get('/api/movements/<int:ident>/suggestions')
    @auth()
    def suggestions(ident):
        m=movement_access(ident); candidates=[]
        for doc in accessible_docs():
            meta=json.loads(doc['metadata']); total=money(meta.get('total'))
            if doc['entity_id']!=m['entity_id'] or total is None or total<=paid_cents(doc['id']) or meta.get('type') not in ('Fatura','Fatura-recibo','Recibo','Nota de crédito'): continue
            if (meta.get('type')=='Nota de crédito')!=(m['amount']>0): continue
            reasons=[]; score=0
            if total-paid_cents(doc['id'])==abs(m['amount']): reasons.append('Montante coincide com o valor em falta'); score+=3
            if meta.get('date'):
                diff=abs((datetime.fromisoformat(meta['date'])-datetime.fromisoformat(m['date'])).days)
                if diff<=30: reasons.append(f'Datas próximas ({diff} dias)'); score+=1
            for key in ('supplier','number'):
                value=str(meta.get(key,'')).strip()
                if len(value)>=3 and value.casefold() in m['description'].casefold(): reasons.append('Fornecedor ou referência na descrição'); score+=2
            if score>=3: candidates.append({'document':doc_out(doc),'reasons':reasons,'score':score})
        return jsonify(sorted(candidates,key=lambda s:s['score'],reverse=True))

    @app.get('/api/requests')
    @auth()
    def requests_list():
        result=[]; allowed={d['id'] for d in accessible_docs()}
        for r in rows('SELECT r.*,u.name AS author FROM requests r JOIN users u ON u.id=r.user_id ORDER BY r.id DESC'):
            if g.user['role']=='admin' or (one('SELECT id FROM grants WHERE user_id=? AND entity_id=?',(g.user['id'],r['entity_id'])) and (r['document_id'] is None or r['document_id'] in allowed)):
                r['replies']=rows('SELECT rr.*,u.name AS author FROM request_replies rr JOIN users u ON u.id=rr.user_id WHERE request_id=? ORDER BY rr.id',(r['id'],)); result.append(r)
        return jsonify(result)
    @app.post('/api/requests')
    @auth()
    def add_request():
        d=data(); eid=int(d.get('entity_id',0)); entity_access(eid); did=d.get('document_id') or None
        if did and doc_access(int(did))['entity_id']!=eid: raise ValueError('Documento de outra entidade.')
        ident=db().execute('INSERT INTO requests(entity_id,document_id,user_id,body,created) VALUES(?,?,?,?,?)',(eid,did,g.user['id'],required(d,'body',5000),now())).lastrowid
        audit('requested','requests',ident); db().commit(); return jsonify(id=ident),201
    @app.post('/api/requests/<int:ident>/reply')
    @auth()
    def reply_request(ident):
        r=one('SELECT * FROM requests WHERE id=?',(ident,))
        if not r: abort(404)
        entity_access(r['entity_id'])
        if r['document_id']: doc_access(r['document_id'])
        d=data(); body=required(d,'body',5000); db().execute('INSERT INTO request_replies(request_id,user_id,body,created) VALUES(?,?,?,?)',(ident,g.user['id'],body,now()))
        if d.get('resolved'): db().execute("UPDATE requests SET status='resolved' WHERE id=?",(ident,))
        audit('replied','requests',ident); db().commit(); return jsonify(ok=True)

    def private_validate(kind,d):
        required(d,'name',200)
        if kind=='works':
            for k in ('budget',):
                if money(d.get(k)) is not None and money(d[k])<0: raise ValueError('Orçamento negativo.')
            for k in ('start','end'):
                if d.get(k): d[k]=date_iso(d[k])
            if d.get('start') and d.get('end') and d['end']<d['start']: raise ValueError('A data final deve ser posterior à inicial.')
            if not isinstance(d.get('categories',[]),list): raise ValueError('Categorias inválidas.')
        if kind=='materials':
            for k in ('price','shipping'):
                if d.get(k) not in ('',None) and money(d[k])<0: raise ValueError('Preço ou portes negativos.')
            for k in ('quantity','coverage','consumption','thickness','width','length','pack_quantity'):
                if d.get(k) not in ('',None):
                    value=Decimal(str(d[k]).replace(',','.'))
                    if not value.is_finite() or value<=0: raise ValueError(f'{k}: indique um número positivo.')
                    d[k]=float(value)
            for k in ('discount','vat'):
                if d.get(k) not in ('',None) and not 0<=Decimal(str(d[k]).replace(',','.'))<=100: raise ValueError('IVA/desconto deve estar entre 0 e 100.')
            if d.get('date'): d['date']=date_iso(d['date'])
            if d.get('document_id'):
                doc=doc_access(int(d['document_id']))
                if d.get('page') and not 1<=int(d['page'])<=int(json.loads(doc['extraction']).get('pages',100)): raise ValueError('Página de origem inválida.')
        if kind=='contacts':
            if d.get('email'): d['email']=email_address(d['email'])
            if d.get('website'):
                u=urlsplit(d['website'])
                if u.scheme not in ('http','https') or not u.hostname: raise ValueError('Website deve começar por https:// ou http://.')
            if not isinstance(d.get('specialties',[]),list): raise ValueError('Especialidades inválidas.')
            if d.get('rating') not in (None,'') and not 0<=int(d['rating'])<=5: raise ValueError('Avaliação entre 0 e 5.')
        if kind=='expenses':
            if d.get('type') not in ('adjudicado','faturado','pago'): raise ValueError('Tipo de registo inválido.')
            if money(d.get('amount')) is None or money(d['amount'])<0: raise ValueError('Indique um valor positivo.')
            d['date']=date_iso(required(d,'date'))
            if not d.get('work_id'): raise ValueError('Selecione uma obra.')
            if d.get('document_id'):
                doc=doc_access(int(d['document_id']))
                if str(doc['work_id'])!=str(d['work_id']): raise ValueError('O comprovativo tem de estar associado à mesma obra.')
                raise ValueError('Os valores dos documentos associados já entram no resumo automaticamente. Use este registo apenas para valores manuais sem documento financeiro.')
        if d.get('work_id'): private_record('works',int(d['work_id']))
        if len(j(d))>30000: raise ValueError('Registo demasiado longo.')
        return d
    @app.get('/api/private/<kind>')
    @auth(True)
    def private_list(kind):
        if kind not in ('works','materials','contacts','expenses'): abort(404)
        result=rows('SELECT * FROM private_records WHERE kind=? ORDER BY id DESC',(kind,))
        for r in result: r['data']=json.loads(r['data'])
        if kind=='works':
            for r in result: r['summary']=work_summary(r)
        if kind=='materials':
            for r in result: r['price_summary']=price_summary(r['data'])
        return jsonify(result)
    @app.post('/api/private/<kind>')
    @auth(True)
    def private_add(kind):
        if kind not in ('works','materials','contacts','expenses'): abort(404)
        d=private_validate(kind,data()); ident=db().execute('INSERT INTO private_records(kind,data,created,updated) VALUES(?,?,?,?)',(kind,j(d),now(),now())).lastrowid
        audit('created',kind,ident); db().commit(); return jsonify(id=ident),201
    @app.get('/api/private/<kind>/<int:ident>')
    @auth(True)
    def private_get(kind,ident): return jsonify(private_record(kind,ident))
    @app.put('/api/private/<kind>/<int:ident>')
    @auth(True)
    def private_update(kind,ident):
        private_record(kind,ident); d=private_validate(kind,data()); db().execute('UPDATE private_records SET data=?,updated=? WHERE id=?',(j(d),now(),ident)); audit('updated',kind,ident); db().commit(); return jsonify(ok=True)
    @app.delete('/api/private/<kind>/<int:ident>')
    @auth(True)
    def private_delete(kind,ident):
        private_record(kind,ident)
        if kind=='works':
            linked=one('SELECT id FROM documents WHERE work_id=?',(ident,)) or any(str(json.loads(r['data']).get('work_id'))==str(ident) for r in rows("SELECT data FROM private_records WHERE kind IN ('materials','expenses')"))
            if linked: raise ValueError('Esta obra tem documentos ou registos associados. Remova as associações primeiro.')
        db().execute('DELETE FROM private_records WHERE id=?',(ident,)); audit('deleted',kind,ident); db().commit(); return jsonify(ok=True)

    def work_summary(work):
        budget=money(work['data'].get('budget')) or 0; awarded=0; invoiced=0; paid=0; unverified=0
        for r in rows("SELECT data FROM private_records WHERE kind='expenses'"):
            d=json.loads(r['data'])
            if str(d.get('work_id'))!=str(work['id']): continue
            amount=money(d.get('amount')) or 0
            if d['type']=='adjudicado': awarded+=amount
            if d['type']=='faturado': invoiced+=amount
            if d['type']=='pago': paid+=amount
        for doc in rows('SELECT * FROM documents WHERE work_id=? AND deleted IS NULL',(work['id'],)):
            m=json.loads(doc['metadata'])
            if m.get('type') not in ('Fatura','Fatura-recibo','Recibo','Nota de crédito'): continue
            # Receipt referencing an already registered invoice must be tagged as comprovativo.
            if m.get('category')=='Comprovativo de pagamento': continue
            sign=-1 if m.get('type')=='Nota de crédito' else 1
            if doc['review']=='validado': invoiced+=sign*(money(m.get('total')) or 0)
            else: unverified+=sign*(money(m.get('total')) or 0)
            paid+=sign*paid_cents(doc['id'])
        return {'budget':budget,'awarded':awarded,'invoiced':invoiced,'paid':paid,'unpaid':max(0,invoiced-paid),'deviation':invoiced-budget,'unverified':unverified,'manual_values_without_receipt':True}

    def price_summary(d):
        price=money(d.get('price'))
        if price is None: return {'missing':['preço'],'unit_net':None}
        p=Decimal(str(d.get('price')).replace(',','.')); discount=Decimal(str(d.get('discount') or 0).replace(',','.')); vat=Decimal(str(d.get('vat') or 0).replace(',','.'))
        net=p*(1-discount/100); gross=net*(1+vat/100)
        coverage=d.get('coverage')
        if not coverage and d.get('width') and d.get('length'): coverage=float(d['width'])*float(d['length'])
        unit=d.get('unit'); divisor=None; label=unit or 'unidade'; missing=[]
        if unit=='m²': divisor=1; label='m²'
        elif coverage: divisor=coverage; label='m²'
        elif unit in ('kg','saco') and d.get('consumption') and d.get('thickness') and (unit=='kg' or d.get('pack_quantity')):
            divisor=(1 if unit=='kg' else float(d['pack_quantity']))/float(d['consumption']); label='m² (estimado)'
        elif unit in ('caixa','saco','rolo','placa'):
            missing=['cobertura por embalagem ou medidas; para argamassa: kg/embalagem, consumo kg/m² e espessura']
        else: divisor=float(d.get('pack_quantity') or 1)
        shipping=money(d.get('shipping')); quantity=Decimal(str(d.get('quantity') or 1)); sp=(Decimal(shipping)/100/quantity) if shipping is not None else None
        vat_known=d.get('vat') not in (None,'')
        return {'original':float(p),'pack_net':round(float(net),4),'pack_gross':round(float(gross),4) if vat_known else None,'unit_net':round(float(net/Decimal(str(divisor))),4) if divisor else None,'unit_gross':round(float(gross/Decimal(str(divisor))),4) if divisor and vat_known else None,'unit':label,'coverage':divisor,'shipping_per_pack':round(float(sp),4) if sp is not None else None,'unit_with_shipping':round(float((gross+sp)/Decimal(str(divisor))),4) if sp is not None and divisor and vat_known else None,'missing':missing,'consumption':d.get('consumption'),'application_thickness':d.get('thickness'),'vat_assumed':not vat_known}
    @app.post('/api/materials/compare')
    @auth(True)
    def compare():
        d=data(); ids=d.get('ids',[])
        if len(ids)<2 or len(ids)>30: raise ValueError('Selecione entre 2 e 30 preços.')
        items=[private_record('materials',int(ident)) for ident in ids]; base=items[0]['data']; result=[]
        area=Decimal(str(d.get('area') or 0).replace(',','.')); waste=Decimal(str(d.get('waste') or 0).replace(',','.'))
        if not area.is_finite() or not waste.is_finite() or area<0 or not 0<=waste<=100: raise ValueError('Área/desperdício inválidos.')
        confirmed=d.get('confirmed',False)
        for item in items:
            m=item['data']; p=price_summary(m)
            exact=bool(base.get('reference') and m.get('reference') and base.get('brand') and m.get('brand') and all(str(base.get(k,'')).casefold()==str(m.get(k,'')).casefold() for k in ('reference','brand','width','length','thickness','unit','coverage','pack_quantity','consumption')))
            potential=bool(base.get('category') and base.get('category')==m.get('category') and p['unit']=='m²')
            status='Produto exatamente igual' if exact else 'Produto potencialmente equivalente' if potential else 'Produto não diretamente comparável'
            estimated='estimado' in p['unit'] or 'estimado' in price_summary(base)['unit']
            same_thickness=bool(m.get('thickness') and base.get('thickness') and str(m['thickness'])==str(base['thickness']))
            accepted=(exact or confirmed) and p['unit_net'] is not None and p['unit']==price_summary(base)['unit'] and (not estimated or same_thickness)
            result.append({'id':item['id'],'data':m,'price':p,'comparison':status,'comparable':accepted,'suggestion_method':'Regras locais; IA externa não configurada'})
            if area>0 and p['coverage'] and p['unit'].startswith('m²'):
                packs=int((area*(1+waste/100)/Decimal(str(p['coverage']))).to_integral_value(rounding=ROUND_CEILING))
                result[-1]['needs']={'packs':packs,'area_with_waste':float(area*(1+waste/100)),'total_net':round(packs*p['pack_net'],2),'total_gross':round(packs*p['pack_gross'],2) if p['pack_gross'] is not None else None,'shipping_known':p['shipping_per_pack'] is not None,'total_with_shipping':round(packs*(p['pack_gross']+p['shipping_per_pack']),2) if p['shipping_per_pack'] is not None and p['pack_gross'] is not None else None}
        comparable=[r for r in result if r['comparable'] and r['price']['unit_gross'] is not None and not r['price']['vat_assumed']]
        best=min(comparable,key=lambda r:r['price']['unit_gross']) if len(comparable)>=2 else None
        for r in result:
            if best and r['comparable'] and not r['price']['vat_assumed']:
                delta=r['price']['unit_gross']-best['price']['unit_gross']; r['difference_eur']=round(delta,4); r['difference_percent']=round(delta/best['price']['unit_gross']*100,2) if best['price']['unit_gross'] else None
        if confirmed:
            audit('comparability_confirmed','materials',detail={'ids':ids}); db().commit()
        return jsonify(items=result,best_id=best['id'] if best else None,warning='Preços históricos dos documentos, sem confirmação de preço ou stock atual. Comparação com IVA; portes separados. Confirme equivalências antes de decidir.')

    def financial_summary(docs):
        months={}; payable=0; review=0; net_total=0; gross_total=0; unverified=0
        for doc in docs:
            m=json.loads(doc['metadata']); total=money(m.get('total')) or 0
            if doc['review']=='por_rever': review+=1; unverified+=total
            if m.get('type') not in ('Fatura','Fatura-recibo','Recibo','Nota de crédito') or m.get('category')=='Comprovativo de pagamento': continue
            sign=-1 if m.get('type')=='Nota de crédito' else 1
            if sign>0 and total>paid_cents(doc['id']): payable+=1
            if doc['review']=='validado':
                month=(m.get('date') or '')[:7]
                if month: months[month]=months.get(month,0)+sign*total
                net_total+=sign*(money(m.get('net')) or 0); gross_total+=sign*total
        return {'review':review,'payable':payable,'months':months,'net':net_total,'gross':gross_total,'unverified':unverified}
    @app.get('/api/dashboard')
    @auth()
    def dashboard():
        docs=[d for d in accessible_docs() if d['visibility']=='financial']; eid=request.args.get('entity_id')
        if eid: entity_access(int(eid)); docs=[d for d in docs if d['entity_id']==int(eid)]
        out=financial_summary(docs); out['documents']=len(docs)
        sql='SELECT m.* FROM movements m WHERE NOT EXISTS(SELECT 1 FROM allocations a WHERE a.movement_id=m.id)'
        unlinked=rows(sql)
        out['unlinked']=sum(1 for m in unlinked if (not eid or m['entity_id']==int(eid)) and (g.user['role']=='admin' or one('SELECT id FROM grants WHERE user_id=? AND entity_id=? AND bank_access=1',(g.user['id'],m['entity_id']))))
        out['requests']=sum(1 for r in requests_list().json if r['status']=='pending' and (not eid or r['entity_id']==int(eid)))
        if g.user['role']=='admin':
            works=rows("SELECT * FROM private_records WHERE kind='works' ORDER BY id DESC")
            for w in works: w['data']=json.loads(w['data']); w['summary']=work_summary(w)
            out['works']=works; out['latest_materials']=private_list('materials').json[:5]
        return jsonify(out)

    @app.get('/api/export')
    @auth()
    def export_period():
        eid=int(request.args.get('entity_id',0)); entity_access(eid); period=request.args.get('period','')
        import re
        if not re.fullmatch(r'\d{4}-\d{2}',period) or not 1<=int(period[5:])<=12: raise ValueError('Selecione o ano e mês.')
        docs=[d for d in accessible_docs() if d['entity_id']==eid and str(json.loads(d['metadata']).get('date') or '')[:7]==period and d['visibility']=='financial']
        allowed={d['id'] for d in docs}; buf=io.BytesIO()
        with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
            stream=io.StringIO(); writer=csv.writer(stream,delimiter=';'); writer.writerow(['ID','Ficheiro','Fornecedor','NIF','Número','Data','Sem IVA EUR','IVA EUR','Total EUR','Pago EUR','Revisão','Contabilidade'])
            for doc in docs:
                m=json.loads(doc['metadata']); z.write(store.get_path(doc['storage'],doc['sha256']),f'originais/{doc["id"]}-{doc["filename"]}')
                writer.writerow([doc['id'],doc['filename'],m.get('supplier',''),m.get('supplier_nif',''),m.get('number',''),m.get('date',''),m.get('net',''),m.get('vat',''),m.get('total',''),paid_cents(doc['id'])/100,doc['review'],doc['accounting']])
            z.writestr('indice.csv',csv_safe(stream.getvalue()).encode('utf-8-sig'))
            grant=one('SELECT bank_access FROM grants WHERE user_id=? AND entity_id=?',(g.user['id'],eid))
            if g.user['role']=='admin' or (grant and grant['bank_access']):
                bank=io.StringIO(); wr=csv.writer(bank,delimiter=';'); wr.writerow(['ID','Banco','Conta','Data','Data valor','Descrição','Montante EUR','Saldo EUR','Classificação'])
                mids=set()
                for m in rows('SELECT * FROM movements WHERE entity_id=? AND date LIKE ?',(eid,period+'%')):
                    mids.add(m['id']); wr.writerow([m['id'],m['bank'],m['account'],m['date'],m['value_date'],m['description'],m['amount']/100,m['balance']/100 if m['balance'] is not None else '',m['classification']])
                z.writestr('movimentos.csv',csv_safe(bank.getvalue()).encode('utf-8-sig'))
                alloc=io.StringIO(); wa=csv.writer(alloc,delimiter=';'); wa.writerow(['Movimento ID','Documento ID','Montante EUR','Data de associação'])
                for a in rows('SELECT * FROM allocations'):
                    if a['document_id'] in allowed: wa.writerow([a['movement_id'],a['document_id'],a['amount']/100,a['created']])
                z.writestr('associacoes.csv',alloc.getvalue().encode('utf-8-sig'))
        audit('period_exported','entities',eid,{'period':period,'documents':len(docs)}); db().commit(); buf.seek(0)
        return send_file(buf,mimetype='application/zip',download_name=f'contabilidade-{eid}-{period}.zip',as_attachment=True)

    @app.get('/api/audit')
    @auth(True)
    def audit_list(): return jsonify(rows('SELECT a.*,u.name AS author FROM audit a LEFT JOIN users u ON u.id=a.user_id ORDER BY a.id DESC LIMIT 300'))
    @app.get('/api/backup')
    @auth(True)
    def backup():
        buf=io.BytesIO()
        with __import__('tempfile').TemporaryDirectory() as tmp:
            name,catalogue=database.export(tmp)
            with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
                z.write(Path(tmp)/name,name)
                for doc in catalogue: z.write(store.get_path(doc['storage'],doc['sha256']),'files/'+doc['storage'])
                z.writestr('LEIA-ME.txt','Cópia privada: contém contas e todos os dados e originais, incluindo documentos apagados. Guarde num local seguro. Restaure com a aplicação parada e uma nova chave de sessão. Não partilhe com o contabilista.')
        audit('backup_created','backup'); db().commit(); buf.seek(0)
        return send_file(buf,mimetype='application/zip',download_name='copia-privada-'+datetime.now().strftime('%Y%m%d')+'.zip',as_attachment=True)
    @app.get('/api/integrations')
    @auth()
    def integrations():
        import shutil
        return jsonify(ocr={'configured':bool(shutil.which('tesseract')),'service':'Tesseract local — sem custo por documento'},pdf={'configured':bool(shutil.which('pdftoppm')),'service':'Poppler / pypdf local'},ai={'configured':False,'service':'Nenhum fornecedor externo configurado; introdução manual e regras locais disponíveis'},storage={'configured':True,'service':'Supabase privado + PostgreSQL; confirmar quotas do plano gratuito' if store.remote else 'Disco privado local + SQLite; custos do alojamento e cópias de segurança'},email={'configured':False,'service':'Convites por link copiado; envio de email não configurado'})

    @app.cli.command('backup')
    def cli_backup():
        """Create an offline-ready private snapshot; schedule externally."""
        import tempfile
        destination=root/'backups'; destination.mkdir(exist_ok=True,mode=0o700)
        target_file=destination/f'{datetime.now().strftime("%Y%m%d-%H%M%S")}.zip'
        with tempfile.TemporaryDirectory() as tmp:
            name,catalogue=database.export(tmp)
            with zipfile.ZipFile(target_file,'w',zipfile.ZIP_DEFLATED) as z:
                z.write(Path(tmp)/name,name)
                for doc in catalogue: z.write(store.get_path(doc['storage'],doc['sha256']),'files/'+doc['storage'])
        os.chmod(target_file,0o600); print(target_file)

    stop=threading.Event(); app.extensions['ocr_stop']=stop
    def worker():
        while not stop.is_set():
            try:
                with app.app_context():
                    g.app_db=app.config['DATABASE']; g.user=None
                    db().execute('BEGIN IMMEDIATE')
                    doc=one("SELECT * FROM documents WHERE ocr_status='queued' AND deleted IS NULL ORDER BY id LIMIT 1")
                    if not doc: db().commit()
                    else:
                        db().execute("UPDATE documents SET ocr_status='processing' WHERE id=?",(doc['id'],)); db().commit()
                        try:
                            result=extract(store.get_path(doc['storage'],doc['sha256']),doc['mime'])
                            db().execute("UPDATE documents SET extraction=?,ocr_status='completed',ocr_error=NULL WHERE id=?",(j(result),doc['id']))
                        except Exception as e:
                            db().execute("UPDATE documents SET ocr_status='error',ocr_error=? WHERE id=?",(str(e)[:1000],doc['id']))
                        audit('ocr_finished','documents',doc['id']); db().commit()
            except Exception:
                app.logger.exception('OCR worker failed')
            stop.wait(1)
    if app.config['START_WORKER']: threading.Thread(target=worker,daemon=True,name='ocr-worker').start()
    return app


def csv_safe(text):
    # Avoid formula injection in downloaded spreadsheets, even for names from OCR.
    source=csv.reader(io.StringIO(text),delimiter=';'); output=io.StringIO(); writer=csv.writer(output,delimiter=';')
    for row in source: writer.writerow(["'"+cell if cell.lstrip().startswith(('=','+','-','@')) else cell for cell in row])
    return output.getvalue()
